"""Build the blank WZ (wydanie z magazynu) print template from the owner's example file.

The owner supplied a hand-filled example (`szablony/wydanie_z_magazynu_szablon.xlsx`). We reuse it
verbatim so fonts/borders/margins/print-setup match his shop 1:1, then:

  * clear the example data rows (7-29) and their per-group merges,
  * normalise the 23-row table grid to a uniform border (C left / H right medium, thin elsewhere),
  * blank the fill-in cells the generator writes at runtime (customer, WZ number, date).

Run once to (re)generate `app/static/templates/wydanie_z_magazynu.xlsx`, which is committed and loaded
by `WorkOrderService.generate_wz_document`. Static chrome (company header, "Odbiorca:", table header,
"Data"/"Odebrał", the stretch-film legal note) stays untouched.
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

SOURCE = Path("szablony/wydanie_z_magazynu_szablon.xlsx")
DEST = Path("app/static/templates/wydanie_z_magazynu.xlsx")

# Data grid the generator fills. Rows 7..29 = 23 line slots; columns C..H.
DATA_FIRST_ROW = 7
DATA_LAST_ROW = 29

THIN = Side(style="thin")
MEDIUM = Side(style="medium")

# Uniform per-column style for the data grid, matching the owner's example row 7 so every one of the 23
# slots renders identically whichever rows the generator fills. Kolor (E) keeps the emphasised 14pt bold
# of the original "FARBA" cell; UWAGI (H) wraps so a "Zam. klienta X (nasze Y)" reference fits.
_GRID_FONT = {
    "C": Font(name="Calibri", size=12),
    "D": Font(name="Calibri", size=12),
    "E": Font(name="Calibri", size=14, bold=True),
    "F": Font(name="Calibri", size=12),
    "G": Font(name="Calibri", size=12),
    "H": Font(name="Calibri", size=12),
}
_GRID_ALIGN = {
    "C": Alignment(horizontal="right", vertical="center"),
    "D": Alignment(horizontal="left", vertical="center"),
    "E": Alignment(horizontal="center", vertical="center", wrap_text=True),
    "F": Alignment(horizontal="center", vertical="center"),
    "G": Alignment(horizontal="right", vertical="center"),
    "H": Alignment(horizontal="center", vertical="center", wrap_text=True),
}


def _grid_border(column: str) -> Border:
    """Uniform table-cell border: outer frame medium on C (left) and H (right), thin everywhere else."""
    left = MEDIUM if column == "C" else THIN
    right = MEDIUM if column == "H" else THIN
    return Border(left=left, right=right, top=THIN, bottom=THIN)


def build() -> None:
    wb = openpyxl.load_workbook(SOURCE)
    ws = wb.active
    ws.title = "WZ"

    # 1. Drop every merged range that lives inside the data grid (per-group Kolor/UWAGI merges from the
    #    example). Footer and header merges (note, Data box, Odbiorca box) are left intact.
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row >= DATA_FIRST_ROW and merged.max_row <= DATA_LAST_ROW:
            ws.unmerge_cells(str(merged))

    # 2. Clear the data grid and give every cell the uniform border/font/alignment.
    for row in range(DATA_FIRST_ROW, DATA_LAST_ROW + 1):
        for column in ("C", "D", "E", "F", "G", "H"):
            cell = ws[f"{column}{row}"]
            cell.value = None
            cell.border = _grid_border(column)
            cell.font = _GRID_FONT[column]
            cell.alignment = _GRID_ALIGN[column]

    # 3. Blank the runtime fill-in cells (kept styled). H1=Odbiorca name, H3=address (handwritten →
    #    stays empty), H4=WZ number, E31=date.
    ws["H1"] = None
    ws["H3"] = None
    ws["H4"] = None
    ws["E31"] = None

    DEST.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DEST)
    print(f"Wrote {DEST}")


if __name__ == "__main__":
    build()
