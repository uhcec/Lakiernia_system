import io
import json
import queue as _queue
import threading
from datetime import UTC, date, datetime
from pathlib import Path
from uuid import uuid4

from fastapi import UploadFile
from pydantic import ValidationError

from app.audit.context import StationContext, ensure_station_allowed
from app.audit.models import AuditEvent
from app.audit.repository import AuditRepository
from app.core.config import get_settings, get_work_order_attachments_path
from app.paints.repository import PaintRepository
from app.paints.service import PaintService
from app.process.enums import ProcessDisplayStatus
from app.process.repository import ProcessRepository
from app.process.service import ProcessService
from app.shared.enums import StationId, WorkOrderStatus, get_work_order_status_label
from app.shared.exceptions import InvalidAttachmentError, InvalidStatusTransitionError, InvalidWorkOrderDataError, ResourceNotFoundError
from app.work_orders.models import WorkOrder, WorkOrderAttachment
from app.work_orders.number_allocator import WorkOrderNumberAllocator
from app.work_orders.repository import WorkOrderRepository
from app.work_orders.schemas import (
    WorkOrderAttachmentRead,
    WorkOrderCreate,
    WorkOrderListResponse,
    WorkOrderQuantityBatchRead,
    WorkOrderRelationSummaryRead,
    WorkOrderRead,
    WorkOrderReleaseUpdate,
    WorkOrderStageStateRead,
    WorkOrderStatusView,
    WorkOrderUpdate,
    calculate_paint_shortage_quantity_kg,
    calculate_predicted_paint_quantity_kg,
    build_work_order_attachment_read,
    build_coating_type_view,
    build_process_status_view,
    build_process_stage_view,
    build_stage_status_info_view,
    build_quantity_batch_type_view,
    build_work_order_relation_summary,
    build_queue_namespace_view,
    build_work_order_kind_view,
    build_production_type_view,
)


class _ExcelPdfConverter:
    """
    Persistent worker thread that keeps one Excel instance alive for the
    lifetime of the process. All conversions are serialized through this
    thread so COM thread-affinity is respected and Excel is never restarted
    between requests.

    The worker is started lazily (see ``_get_excel_pdf_converter``) so importing
    this module has no side effects and does not require Windows/Excel. If Excel
    or ``pywin32`` is unavailable the first conversion fails fast with a clear
    error instead of hanging until the timeout.
    """

    def __init__(self) -> None:
        self._q: _queue.Queue = _queue.Queue()
        self._ready = threading.Event()
        self._init_error: Exception | None = None
        t = threading.Thread(target=self._worker, daemon=True, name="excel-pdf-worker")
        t.start()

    def _worker(self) -> None:
        try:
            import pythoncom
            import win32com.client

            pythoncom.CoInitialize()
            excel = win32com.client.Dispatch("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.AskToUpdateLinks = False
        except Exception as exc:  # pywin32/Excel missing or COM init failed
            self._init_error = exc
            self._ready.set()
            return

        self._ready.set()
        while True:
            task = self._q.get()
            if task is None:
                break
            xlsx_str, pdf_str, done_event, err_box = task
            try:
                wb = excel.Workbooks.Open(xlsx_str, UpdateLinks=0, ReadOnly=0)
                ws = wb.Worksheets(1)
                ws.Range("D6").Font.Size = 16
                ws.Range("D6").Font.Bold = True
                ps = ws.PageSetup
                ps.PaperSize = 9  # A4
                ps.Orientation = 1  # portrait
                ps.CenterHorizontally = True
                ps.CenterVertically = True
                ps.Zoom = False
                ps.FitToPagesWide = 1
                ps.FitToPagesTall = 1
                wb.ExportAsFixedFormat(0, pdf_str)
                wb.Close(False)
            except Exception as exc:
                err_box.append(exc)
            finally:
                done_event.set()

    def convert(self, xlsx_path: Path, pdf_path: Path, timeout: float = 60.0) -> None:
        # Wait for the worker to finish initializing Excel before queueing work.
        if not self._ready.wait(timeout=timeout):
            raise RuntimeError("Konwersja PDF przekroczyła limit czasu (Excel nie odpowiada).")
        if self._init_error is not None:
            raise RuntimeError(
                "Generowanie dokumentu wymaga systemu Windows z zainstalowanym pakietem Excel i biblioteką pywin32."
            ) from self._init_error

        done_event = threading.Event()
        err_box: list[Exception] = []
        self._q.put((str(xlsx_path.resolve()), str(pdf_path.resolve()), done_event, err_box))
        if not done_event.wait(timeout=timeout):
            raise RuntimeError("Konwersja PDF przekroczyła limit czasu.")
        if err_box:
            raise err_box[0]


_excel_pdf_converter: _ExcelPdfConverter | None = None
_excel_pdf_converter_lock = threading.Lock()


def _get_excel_pdf_converter() -> _ExcelPdfConverter:
    global _excel_pdf_converter
    if _excel_pdf_converter is None:
        with _excel_pdf_converter_lock:
            if _excel_pdf_converter is None:
                _excel_pdf_converter = _ExcelPdfConverter()
    return _excel_pdf_converter


class WorkOrderService:
    ALLOWED_ATTACHMENT_CONTENT_TYPES = {"application/pdf"}
    PDF_MAGIC_BYTES = b"%PDF-"

    def __init__(self, repository: WorkOrderRepository) -> None:
        self.repository = repository
        self.process_service = ProcessService(ProcessRepository(repository.session))
        self.paint_service = PaintService(PaintRepository(repository.session))
        self.number_allocator = WorkOrderNumberAllocator(repository.session)
        self.audit_repository = AuditRepository(repository.session)

    @staticmethod
    def derive_status(process_status: ProcessDisplayStatus) -> WorkOrderStatus:
        if process_status == ProcessDisplayStatus.INCOMPLETE:
            return WorkOrderStatus.INCOMPLETE
        if process_status == ProcessDisplayStatus.TO_BE_PLANNED:
            return WorkOrderStatus.ENTERED
        if process_status == ProcessDisplayStatus.READY:
            return WorkOrderStatus.FINISHED
        return WorkOrderStatus.IN_PROGRESS

    @classmethod
    def build_status_view(cls, process_status: ProcessDisplayStatus) -> WorkOrderStatusView:
        status = cls.derive_status(process_status)
        return WorkOrderStatusView(value=status, label=get_work_order_status_label(status))

    def _get_current_paint_stock_quantity_kg(self, work_order: WorkOrder) -> float | None:
        if not work_order.ral_code or not work_order.paint_producer or not work_order.paint_code:
            return None
        return self.paint_service.lookup_current_stock(
            ral_code=work_order.ral_code,
            paint_producer=work_order.paint_producer,
            paint_code=work_order.paint_code,
        ).current_quantity_kg

    def create_work_order(self, payload: WorkOrderCreate, station_context: StationContext | None = None) -> WorkOrderRead:
        self._ensure_office_station_action(station_context, "office work-order actions")
        order_number = self.number_allocator.allocate()
        paint, paint_stock_snapshot_quantity_kg = self.paint_service.capture_work_order_snapshot(
            ral_code=payload.ral_code,
            paint_producer=payload.paint_producer,
            paint_code=payload.paint_code,
        )

        work_order = WorkOrder(
            paint_id=paint.id,
            order_number=order_number,
            external_order_number=payload.external_order_number,
            title=payload.title or order_number,
            total_quantity=payload.total_quantity,
            description=payload.description,
            customer_name=payload.customer_name,
            ral_code=payload.ral_code,
            paint_code=payload.paint_code,
            paint_producer=payload.paint_producer,
            paint_stock_snapshot_quantity_kg=paint_stock_snapshot_quantity_kg,
            consumed_paint_quantity_kg=0,
            coating_type=payload.coating_type,
            work_order_kind=payload.work_order_kind,
            is_ocynk=payload.is_ocynk,
            received_date=payload.received_date,
            planned_pickup_date=payload.planned_pickup_date,
            total_area_m2=payload.total_area_m2,
            is_incomplete_on_arrival=payload.is_incomplete_on_arrival,
            profiles_quantity=payload.profiles_quantity,
            profiles_to_count=payload.profiles_to_count,
            accessories_quantity=payload.accessories_quantity,
            accessories_to_count=payload.accessories_to_count,
            notes=payload.notes,
        )

        try:
            persisted_work_order = self.repository.add(work_order)
            self.process_service.initialize_quantity_tracking(persisted_work_order)
            self.process_service.initialize_stage_states(persisted_work_order)
            self._record_work_order_created(persisted_work_order, station_context)
            self.repository.commit()
        except Exception:
            self.repository.rollback()
            raise

        return self._to_read_model(persisted_work_order)

    def get_work_order(self, work_order_id: int) -> WorkOrderRead:
        work_order = self._get_existing_work_order(work_order_id)
        return self._to_read_model(work_order)

    def update_work_order(
        self,
        work_order_id: int,
        payload: WorkOrderUpdate,
        station_context: StationContext | None = None,
    ) -> WorkOrderRead:
        ensure_station_allowed(
            station_context,
            allowed_station_ids={StationId.OFFICE, StationId.MANAGER},
            action="work-order update",
        )
        work_order = self._get_existing_work_order(work_order_id)
        updates = payload.model_dump(exclude_unset=True)
        validated_payload = self._validate_work_order_state(work_order, updates)
        paint_identity_changed = any(field_name in updates for field_name in {"ral_code", "paint_code", "paint_producer"})
        resolved_paint = None
        resolved_snapshot = None
        if paint_identity_changed:
            resolved_paint, resolved_snapshot = self.paint_service.capture_work_order_snapshot(
                ral_code=validated_payload.ral_code,
                paint_producer=validated_payload.paint_producer,
                paint_code=validated_payload.paint_code,
            )
        changed_fields = {
            field_name: {
                "before": getattr(work_order, field_name),
                "after": getattr(validated_payload, field_name),
            }
            for field_name in updates
            if getattr(work_order, field_name) != getattr(validated_payload, field_name)
        }

        for field_name in updates:
            setattr(work_order, field_name, getattr(validated_payload, field_name))

        if resolved_paint is not None and resolved_snapshot is not None:
            if work_order.paint_id != resolved_paint.id:
                changed_fields["paint_id"] = {"before": work_order.paint_id, "after": resolved_paint.id}
            if work_order.paint_stock_snapshot_quantity_kg != resolved_snapshot:
                changed_fields["paint_stock_snapshot_quantity_kg"] = {
                    "before": work_order.paint_stock_snapshot_quantity_kg,
                    "after": resolved_snapshot,
                }
            work_order.paint_id = resolved_paint.id
            work_order.paint_stock_snapshot_quantity_kg = resolved_snapshot

        work_order.updated_at = datetime.now(UTC)
        if changed_fields:
            self._record_audit_event(
                event_type="work_order.updated",
                station_context=station_context,
                work_order=work_order,
                target_type="work_order",
                target_id=work_order.id,
                summary=f"Updated work order {work_order.order_number}.",
                details={
                    "updated_fields": sorted(changed_fields),
                    "changes": self._json_ready(changed_fields),
                },
            )

        try:
            self.repository.commit()
        except Exception:
            self.repository.rollback()
            raise

        return self._to_read_model(work_order)

    def _validate_work_order_state(self, work_order: WorkOrder, overrides: dict[str, object] | None = None) -> WorkOrderCreate:
        merged_payload = {
            "title": work_order.title,
            "external_order_number": work_order.external_order_number,
            "total_quantity": work_order.total_quantity,
            "description": work_order.description,
            "customer_name": work_order.customer_name,
            "ral_code": work_order.ral_code,
            "paint_code": work_order.paint_code,
            "paint_producer": work_order.paint_producer,
            "coating_type": work_order.coating_type,
            "work_order_kind": work_order.work_order_kind,
            "is_ocynk": work_order.is_ocynk,
            "received_date": work_order.received_date,
            "planned_pickup_date": work_order.planned_pickup_date,
            "total_area_m2": work_order.total_area_m2,
            "is_incomplete_on_arrival": work_order.is_incomplete_on_arrival,
            "profiles_quantity": work_order.profiles_quantity,
            "profiles_to_count": work_order.profiles_to_count,
            "accessories_quantity": work_order.accessories_quantity,
            "accessories_to_count": work_order.accessories_to_count,
            "notes": work_order.notes,
        }
        if overrides:
            merged_payload.update(overrides)

        try:
            return WorkOrderCreate.model_validate(merged_payload)
        except ValidationError as exc:
            first_error = exc.errors()[0]
            raise InvalidWorkOrderDataError(str(first_error.get("msg", "Nieprawidłowe dane zlecenia."))) from exc

    def list_work_orders(
        self,
        *,
        limit: int,
        offset: int,
        process_status: ProcessDisplayStatus | None = None,
        external_order_number: str | None = None,
    ) -> WorkOrderListResponse:
        external_order_number_query = self._normalize_optional_search(external_order_number)
        repository_total = self.repository.count(external_order_number_query=external_order_number_query)
        all_items = self.repository.list(
            limit=repository_total,
            offset=0,
            external_order_number_query=external_order_number_query,
        )
        # The office list shows client work orders only. Correction sub-orders ("P...") are internal
        # production records; their open state is surfaced via the "N" marker on the parent and the
        # parent's correction context, so they are not listed as separate office rows.
        filtered_items = [item for item in all_items if not item.is_correction]
        if process_status is not None:
            filtered_items = [
                item for item in filtered_items if self.process_service.get_display_status(item) == process_status
            ]
        total = len(filtered_items)
        items = filtered_items[offset : offset + limit]

        return WorkOrderListResponse(
            items=[self._to_read_model(item) for item in items],
            total=total,
            limit=limit,
            offset=offset,
        )

    @staticmethod
    def _normalize_optional_search(value: str | None) -> str | None:
        if value is None:
            return None
        normalized_value = value.strip()
        return normalized_value or None

    def reject_manual_status_transition(self, work_order_id: int) -> WorkOrderRead:
        self._get_existing_work_order(work_order_id)
        raise InvalidStatusTransitionError(
            "Work order status is derived from process status and cannot be changed manually."
        )

    def mark_complete_arrival(
        self,
        work_order_id: int,
        station_context: StationContext | None = None,
    ) -> WorkOrderRead:
        self._ensure_office_station_action(station_context, "office work-order actions")
        work_order = self._get_existing_work_order(work_order_id)

        if not self._has_incomplete_marker(work_order):
            raise InvalidStatusTransitionError(
                f"Work order {work_order_id} is not marked as incomplete on arrival."
            )

        self._validate_work_order_state(work_order, {"is_incomplete_on_arrival": False})

        work_order.is_incomplete_on_arrival = False
        work_order.updated_at = datetime.now(UTC)
        self._record_audit_event(
            event_type="work_order.arrival_completed",
            station_context=station_context,
            work_order=work_order,
            target_type="work_order",
            target_id=work_order.id,
            summary=f"Marked work order {work_order.order_number} as complete on arrival.",
            details={
                "is_incomplete_on_arrival": False,
            },
        )

        try:
            self.repository.commit()
        except Exception:
            self.repository.rollback()
            raise

        return self._to_read_model(work_order)

    def update_release_status(
        self,
        work_order_id: int,
        payload: WorkOrderReleaseUpdate,
        station_context: StationContext | None = None,
    ) -> WorkOrderRead:
        self._ensure_release_station_action(station_context)
        work_order = self._get_existing_work_order(work_order_id)
        process_status = self.process_service.get_display_status(work_order)
        if process_status != ProcessDisplayStatus.READY:
            raise InvalidStatusTransitionError("Only READY work orders can be marked as released to client.")

        current_release_state = work_order.released_to_client_at is not None
        if current_release_state == payload.is_released_to_client:
            return self._to_read_model(work_order)

        now = datetime.now(UTC)
        work_order.released_to_client_at = now if payload.is_released_to_client else None
        work_order.updated_at = now
        self._record_audit_event(
            event_type="work_order.release_status_updated",
            station_context=station_context,
            work_order=work_order,
            target_type="work_order",
            target_id=work_order.id,
            summary=f"Updated release status for work order {work_order.order_number}.",
            details={
                "is_released_to_client": payload.is_released_to_client,
                "released_to_client_at": work_order.released_to_client_at,
            },
        )

        try:
            self.repository.commit()
        except Exception:
            self.repository.rollback()
            raise

        return self._to_read_model(work_order)

    @staticmethod
    def _has_incomplete_marker(work_order: WorkOrder) -> bool:
        return work_order.is_incomplete_on_arrival

    async def upload_attachment(
        self,
        work_order_id: int,
        file: UploadFile,
        *,
        attachment_type: str | None = None,
        description: str | None = None,
        station_context: StationContext | None = None,
    ) -> WorkOrderAttachmentRead:
        self._ensure_office_station_action(station_context, "office attachment actions")
        work_order = self._get_existing_work_order(work_order_id)
        original_filename = Path(file.filename or "").name.strip()
        if not original_filename:
            raise InvalidAttachmentError("Plik musi mieć nazwę.")
        if file.content_type not in self.ALLOWED_ATTACHMENT_CONTENT_TYPES:
            raise InvalidAttachmentError("Dozwolone są tylko pliki PDF.", status_code=415)

        settings = get_settings()
        root_path = get_work_order_attachments_path()
        attachment_directory = root_path / str(work_order.id)
        attachment_directory.mkdir(parents=True, exist_ok=True)

        stored_filename = f"{uuid4().hex}.pdf"
        destination_path = attachment_directory / stored_filename
        relative_storage_path = destination_path.relative_to(root_path).as_posix()

        file_size = 0
        try:
            with destination_path.open("wb") as output_file:
                header = await file.read(len(self.PDF_MAGIC_BYTES))
                if not header:
                    raise InvalidAttachmentError("Nie można przesłać pustego pliku PDF.")
                if header != self.PDF_MAGIC_BYTES:
                    raise InvalidAttachmentError("Przesłany plik nie jest prawidłowym PDF.")

                file_size += len(header)
                output_file.write(header)

                while chunk := await file.read(1024 * 1024):
                    file_size += len(chunk)
                    if file_size > settings.work_order_attachment_max_size_bytes:
                        raise InvalidAttachmentError(
                            f"Plik przekracza limit {settings.work_order_attachment_max_size_bytes} B.",
                            status_code=413,
                        )
                    output_file.write(chunk)

            attachment = WorkOrderAttachment(
                work_order_id=work_order.id,
                original_filename=original_filename,
                stored_filename=stored_filename,
                storage_path=relative_storage_path,
                content_type=file.content_type,
                file_size=file_size,
                attachment_type=self._normalize_optional_text(attachment_type),
                description=self._normalize_optional_text(description),
            )

            persisted_attachment = self.repository.add_attachment(attachment)
            self._record_attachment_audit_event(
                event_type="work_order.attachment_uploaded",
                station_context=station_context,
                work_order=work_order,
                attachment=persisted_attachment,
                summary=f"Uploaded attachment {persisted_attachment.original_filename} for work order {work_order.order_number}.",
            )
            self.repository.commit()
        except Exception:
            self.repository.rollback()
            if destination_path.exists():
                destination_path.unlink()
            raise
        finally:
            await file.close()

        return build_work_order_attachment_read(persisted_attachment)

    def list_attachments(self, work_order_id: int) -> list[WorkOrderAttachmentRead]:
        self._get_existing_work_order(work_order_id)
        return [
            build_work_order_attachment_read(attachment)
            for attachment in self.repository.list_attachments(work_order_id)
        ]

    def get_attachment_download(self, work_order_id: int, attachment_id: int) -> tuple[WorkOrderAttachmentRead, Path]:
        attachment = self._get_existing_attachment(work_order_id, attachment_id)
        absolute_path = self._resolve_attachment_storage_path(attachment, require_exists=True)
        assert absolute_path is not None
        return build_work_order_attachment_read(attachment), absolute_path

    def generate_quality_control_document(self, work_order_id: int) -> tuple[io.BytesIO, str]:
        import tempfile
        import openpyxl

        work_order = self._get_existing_work_order(work_order_id)

        paint_parts = [p for p in [work_order.ral_code, work_order.paint_producer, work_order.paint_code] if p]
        paint_description = " / ".join(paint_parts) if paint_parts else "-"

        template_path = Path(__file__).parent.parent / "static" / "templates" / "karta_kontroli_jakosci.xlsx"

        wb = openpyxl.load_workbook(template_path)
        ws = wb.active
        ws["A4"] = work_order.order_number
        ws["D5"] = work_order.customer_name
        ws["D6"] = paint_description

        with tempfile.TemporaryDirectory() as tmp_dir:
            xlsx_path = Path(tmp_dir) / f"KJ_{work_order.order_number}.xlsx"
            pdf_path = Path(tmp_dir) / f"KJ_{work_order.order_number}.pdf"
            wb.save(xlsx_path)
            _get_excel_pdf_converter().convert(xlsx_path, pdf_path)
            pdf_bytes = pdf_path.read_bytes()

        return io.BytesIO(pdf_bytes), f"KJ_{work_order.order_number}.pdf"

    def delete_attachment(
        self,
        work_order_id: int,
        attachment_id: int,
        station_context: StationContext | None = None,
    ) -> None:
        self._ensure_office_station_action(station_context, "office attachment actions")
        attachment = self._get_existing_attachment(work_order_id, attachment_id)
        work_order = self._get_existing_work_order(work_order_id)
        file_path = self._resolve_attachment_storage_path(attachment, require_exists=False)

        try:
            self._record_attachment_audit_event(
                event_type="work_order.attachment_deleted",
                station_context=station_context,
                work_order=work_order,
                attachment=attachment,
                summary=f"Deleted attachment {attachment.original_filename} from work order {work_order.order_number}.",
            )
            self.repository.delete_attachment(attachment)
            self.repository.commit()
        except Exception:
            self.repository.rollback()
            raise

        if file_path is not None and file_path.exists():
            file_path.unlink()

        parent_directory = file_path.parent if file_path is not None else None
        if parent_directory is not None and parent_directory.exists() and not any(parent_directory.iterdir()):
            parent_directory.rmdir()

    @staticmethod
    def _ensure_office_station_action(station_context: StationContext | None, action: str) -> None:
        ensure_station_allowed(
            station_context,
            allowed_station_ids={StationId.OFFICE},
            action=action,
        )

    @staticmethod
    def _ensure_release_station_action(station_context: StationContext | None) -> None:
        ensure_station_allowed(
            station_context,
            allowed_station_ids={StationId.OFFICE},
            action="work-order release actions",
        )

    @staticmethod
    def _resolve_attachment_storage_path(attachment: WorkOrderAttachment, *, require_exists: bool) -> Path | None:
        root_path = get_work_order_attachments_path()
        absolute_path = (root_path / attachment.storage_path).resolve()
        try:
            absolute_path.relative_to(root_path)
        except ValueError:
            if require_exists:
                raise ResourceNotFoundError(f"Attachment {attachment.id} file not found.")
            return None

        if require_exists and not absolute_path.exists():
            raise ResourceNotFoundError(f"Attachment {attachment.id} file not found.")
        return absolute_path

    def _get_existing_work_order(self, work_order_id: int) -> WorkOrder:
        work_order = self.repository.get_by_id(work_order_id)
        if work_order is None:
            raise ResourceNotFoundError(f"Work order {work_order_id} not found.")
        return work_order

    def _get_existing_attachment(self, work_order_id: int, attachment_id: int) -> WorkOrderAttachment:
        self._get_existing_work_order(work_order_id)
        attachment = self.repository.get_attachment(work_order_id, attachment_id)
        if attachment is None:
            raise ResourceNotFoundError(f"Attachment {attachment_id} not found for work order {work_order_id}.")
        return attachment

    @staticmethod
    def _normalize_optional_text(value: str | None) -> str | None:
        if value is None:
            return None
        normalized = value.strip()
        return normalized or None

    def _record_work_order_created(self, work_order: WorkOrder, station_context: StationContext | None) -> None:
        self._record_audit_event(
            event_type="work_order.created",
            station_context=station_context,
            work_order=work_order,
            target_type="work_order",
            target_id=work_order.id,
            summary=f"Created work order {work_order.order_number}.",
            details={
                "order_number": work_order.order_number,
                "external_order_number": work_order.external_order_number,
                "work_order_kind": work_order.work_order_kind,
                "is_ocynk": work_order.is_ocynk,
                "is_incomplete_on_arrival": work_order.is_incomplete_on_arrival,
            },
        )

    def _record_attachment_audit_event(
        self,
        *,
        event_type: str,
        station_context: StationContext | None,
        work_order: WorkOrder,
        attachment: WorkOrderAttachment,
        summary: str,
    ) -> None:
        self._record_audit_event(
            event_type=event_type,
            station_context=station_context,
            work_order=work_order,
            target_type="work_order_attachment",
            target_id=attachment.id,
            summary=summary,
            details={
                "attachment_id": attachment.id,
                "original_filename": attachment.original_filename,
                "stored_filename": attachment.stored_filename,
                "content_type": attachment.content_type,
                "file_size": attachment.file_size,
                "attachment_type": attachment.attachment_type,
            },
        )

    def _record_audit_event(
        self,
        *,
        event_type: str,
        station_context: StationContext | None,
        work_order: WorkOrder,
        target_type: str,
        target_id: int | None,
        summary: str,
        details: dict[str, object],
    ) -> None:
        station_id = station_context.station_id if station_context is not None else StationId.SYSTEM
        self.audit_repository.add(
            AuditEvent(
                event_type=event_type,
                station_id=station_id,
                work_order_id=work_order.id,
                target_type=target_type,
                target_id=target_id,
                summary=summary,
                details_json=json.dumps(self._json_ready(details), separators=(",", ":")),
            )
        )

    @classmethod
    def _json_ready(cls, value: object) -> object:
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        if hasattr(value, "value"):
            return getattr(value, "value")
        if isinstance(value, dict):
            return {key: cls._json_ready(nested_value) for key, nested_value in value.items()}
        if isinstance(value, list):
            return [cls._json_ready(nested_value) for nested_value in value]
        return value

    def _to_read_model(self, work_order: WorkOrder) -> WorkOrderRead:
        waiting_ok_quantity = self.process_service.get_waiting_ok_quantity(work_order)
        pending_correction_quantity = self.process_service.get_pending_correction_quantity(work_order)
        total_progress_quantity = self.process_service.get_total_progress_quantity(work_order)
        process_status = self.process_service.get_display_status(work_order)

        # Overdue is business meaning, so it is owned by the backend (NA44), not computed on the
        # renderer: the planned pickup day has passed and the order is not yet READY. Date-only,
        # server-local — strictly before today (the pickup day itself is not yet "past the deadline").
        is_overdue = (
            work_order.planned_pickup_date is not None
            and process_status != ProcessDisplayStatus.READY
            and work_order.planned_pickup_date < date.today()
        )

        predicted_paint_quantity_kg = calculate_predicted_paint_quantity_kg(work_order.total_area_m2)
        current_paint_stock_quantity_kg = self._get_current_paint_stock_quantity_kg(work_order)
        paint_inventory_shortage_quantity_kg = calculate_paint_shortage_quantity_kg(
            predicted_paint_quantity_kg,
            current_paint_stock_quantity_kg,
        )

        return WorkOrderRead(
            id=work_order.id,
            order_number=work_order.order_number,
            external_order_number=work_order.external_order_number,
            title=work_order.title,
            total_quantity=work_order.total_quantity,
            completed_quantity=work_order.completed_quantity,
            current_active_quantity=work_order.current_active_quantity,
            waiting_ok_quantity=waiting_ok_quantity,
            pending_correction_quantity=pending_correction_quantity,
            total_progress_quantity=total_progress_quantity,
            is_correction=work_order.is_correction,
            is_historical_correction=self.process_service.is_historical_correction(work_order),
            is_decoral_incomplete=self.process_service.is_decoral_incomplete(work_order),
            work_order_kind=build_work_order_kind_view(work_order.work_order_kind),
            is_ocynk=work_order.is_ocynk,
            correction_reason_category=work_order.correction_reason_category,
            correction_reason_note=work_order.correction_reason_note,
            parent_order=(
                self._build_relation_summary(work_order.parent_order)
                if work_order.parent_order is not None
                else None
            ),
            correction_orders=[
                self._build_relation_summary(correction_order)
                for correction_order in self._get_ordered_correction_orders(work_order)
            ],
            description=work_order.description,
            customer_name=work_order.customer_name,
            ral_code=work_order.ral_code,
            paint_code=work_order.paint_code,
            paint_producer=work_order.paint_producer,
            predicted_paint_quantity_kg=predicted_paint_quantity_kg,
            current_paint_stock_quantity_kg=current_paint_stock_quantity_kg,
            paint_inventory_insufficient_for_prediction=paint_inventory_shortage_quantity_kg is not None,
            paint_inventory_shortage_quantity_kg=paint_inventory_shortage_quantity_kg,
            paint_stock_snapshot_quantity_kg=work_order.paint_stock_snapshot_quantity_kg,
            consumed_paint_quantity_kg=self.paint_service.get_consumed_paint_kg_for_work_orders(work_order.correction_chain_ids),
            status=self.build_status_view(process_status),
            process_status=build_process_status_view(process_status),
            coating_type=build_coating_type_view(work_order.coating_type),
            production_type=(
                build_production_type_view(work_order.production_type)
                if work_order.production_type is not None
                else None
            ),
            queue_namespace=build_queue_namespace_view(work_order.queue_namespace),
            queue_position=work_order.queue_position,
            stage_states=[
                WorkOrderStageStateRead(
                    stage=build_process_stage_view(stage_state.stage),
                    status=build_stage_status_info_view(stage_state.status),
                )
                for stage_state in self._get_ordered_stage_states(work_order)
            ],
            batches=[
                WorkOrderQuantityBatchRead(
                    id=batch.id,
                    quantity=batch.quantity,
                    type=build_quantity_batch_type_view(batch.type),
                    parent_batch_id=batch.parent_batch_id,
                    is_active=batch.is_active,
                    is_finished=not batch.is_active,
                    current_stage=(build_process_stage_view(batch.current_stage) if batch.current_stage is not None else None),
                    current_stage_status=(
                        build_stage_status_info_view(batch.current_stage_status)
                        if batch.current_stage_status is not None
                        else None
                    ),
                    is_partial=batch.is_partial,
                    created_at=batch.created_at,
                )
                for batch in self._get_ordered_batches(work_order)
            ],
            received_date=work_order.received_date,
            planned_pickup_date=work_order.planned_pickup_date,
            is_overdue=is_overdue,
            is_released_to_client=work_order.released_to_client_at is not None,
            released_to_client_at=work_order.released_to_client_at,
            total_area_m2=work_order.total_area_m2,
            is_incomplete_on_arrival=work_order.is_incomplete_on_arrival,
            surface_preparation_counted_quantity=self.process_service.get_surface_preparation_counted_quantity(work_order),
            surface_preparation_quantity_difference=self.process_service.get_surface_preparation_quantity_difference(work_order),
            profiles_quantity=work_order.profiles_quantity,
            profiles_to_count=work_order.profiles_to_count,
            accessories_quantity=work_order.accessories_quantity,
            accessories_to_count=work_order.accessories_to_count,
            notes=work_order.notes,
            created_at=work_order.created_at,
            updated_at=work_order.updated_at,
        )

    def _get_ordered_stage_states(self, work_order: WorkOrder):
        stage_order = {
            stage: index
            for index, stage in enumerate(self.process_service.get_stage_sequence_for_work_order(work_order))
        }
        return sorted(
            [stage_state for stage_state in work_order.stage_states if stage_state.stage in stage_order],
            key=lambda stage_state: stage_order[stage_state.stage],
        )

    @staticmethod
    def _get_ordered_batches(work_order: WorkOrder):
        return sorted(work_order.quantity_batches, key=lambda batch: (batch.created_at, batch.id))

    @staticmethod
    def _get_ordered_correction_orders(work_order: WorkOrder) -> list[WorkOrder]:
        return sorted(work_order.correction_orders, key=lambda correction_order: (correction_order.created_at, correction_order.id))

    def _build_relation_summary(self, work_order: WorkOrder):
        summary = build_work_order_relation_summary(work_order)
        summary.process_status = build_process_status_view(self.process_service.get_display_status(work_order))
        summary.is_historical_correction = self.process_service.is_historical_correction(work_order)
        summary.is_decoral_incomplete = self.process_service.is_decoral_incomplete(work_order)
        return summary