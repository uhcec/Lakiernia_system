import openpyxl
from openpyxl.styles import Alignment, Font
import io
from pathlib import Path

template_path = Path("app/static/templates/karta_kontroli_jakosci.xlsx")
wb = openpyxl.load_workbook(template_path)
ws = wb.active

print("Arkusz:", ws.title)
print("Wymiary:", ws.dimensions)
print("A4:", ws["A4"].value)
print("D5:", ws["D5"].value)
print("D6:", ws["D6"].value)

ws["A4"] = "0001"
ws["D5"] = "Testowy Klient Sp. z o.o."
ws["D6"] = "RAL 9005 / Farby Polskie / KOD-TEST-001"
ws["D6"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws["D6"].font = Font(name="Calibri", bold=True, size=16)

buffer = io.BytesIO()
wb.save(buffer)
buffer.seek(0)

# Zapisz do pliku i sprawdź
out = Path("scripts/test_output.xlsx")
out.write_bytes(buffer.read())
size = out.stat().st_size
print(f"\nZapisany plik: {size} bajtów")

# Wczytaj z powrotem i sprawdź
wb2 = openpyxl.load_workbook(out)
ws2 = wb2.active
print("A4:", ws2["A4"].value)
print("D5:", ws2["D5"].value)
print("D6:", ws2["D6"].value)
print("D6 font size:", ws2["D6"].font.sz)
